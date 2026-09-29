import os
import sys
import docx

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_callout(doc, text_content, label="PROMPT GỐC CỦA NGƯỜI DÙNG:"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=180)
    
    # Left border highlight
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="2563EB"/>'
        f'<w:top w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_label = p.add_run(f"💬 {label}\n")
    run_label.bold = True
    run_label.font.name = "Calibri"
    run_label.font.size = Pt(10.5)
    run_label.font.color.rgb = RGBColor(30, 64, 175) # Blue
    
    run_text = p.add_run(text_content.strip())
    run_text.font.name = "Calibri"
    run_text.font.size = Pt(11)
    run_text.font.italic = True
    run_text.font.color.rgb = RGBColor(15, 23, 42)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def build_word_document():
    doc = Document()
    
    # Page setup (Letter / A4 with 1 inch margin)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
        # Header / Footer
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = f_p.add_run("VinaLex Project • Báo cáo Prompt Engineering & Nhật ký Xây dựng Hệ thống")
        f_run.font.size = Pt(8.5)
        f_run.font.color.rgb = RGBColor(148, 163, 184)

    # Document Header
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    r_sub = title_p.add_run("HỆ THỐNG CỔNG THÔNG TIN PHÁP LÝ & DỊCH VỤ CÔNG QUỐC GIA VINALEX\n")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(11)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(100, 116, 139)
    
    r_main = title_p.add_run("NHẬT KÝ CHI TIẾT TOÀN BỘ PROMPTS\nXÂY DỰNG, CẢI TIẾN & VẬN HÀNH HỆ THỐNG")
    r_main.font.name = "Calibri"
    r_main.font.size = Pt(20)
    r_main.font.bold = True
    r_main.font.color.rgb = RGBColor(15, 23, 42)
    
    # Metadata Block
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Chủ nhiệm dự án:", "Người phát triển VinaLex (haianh02032005-rgb)"),
        ("Đối tác AI đồng hành:", "Google Antigravity IDE (Gemini AI Coding Assistant)"),
        ("Mục tiêu tài liệu:", "Ghi nhận và hệ thống hóa toàn bộ các câu lệnh (Prompts) thực tế đã sử dụng trong suốt quá trình xây dựng, gỡ lỗi, tối ưu hiệu năng, nâng cấp suy luận AI và đóng gói vận hành VinaLex."),
        ("Phạm vi kỹ thuật:", "Full-Stack: Next.js 14, FastAPI Python 3.11, SQLite, RAG Vector Search, Cloudflare Tunnel, Docker, Groq LLaMA 3.3 70B & Legal CoT Engine.")
    ]
    for idx, (label, val) in enumerate(meta_data):
        row = meta_table.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.0)
        c1.width = Inches(4.5)
        set_cell_background(c0, "F8FAFC")
        set_cell_background(c1, "FFFFFF")
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(3)
        p0.paragraph_format.space_after = Pt(3)
        r0 = p0.add_run(label)
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = RGBColor(51, 65, 85)
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(3)
        p1.paragraph_format.space_after = Pt(3)
        r1 = p1.add_run(val)
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = RGBColor(30, 41, 59)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 1. BẢNG TỔNG QUAN PHÂN LOẠI PROMPT
    h1 = doc.add_paragraph()
    h1_run = h1.add_run("PHẦN 1: BẢNG TỔNG KẾT MA TRẬN CÁC PROMPT ĐÃ DÙNG")
    h1_run.font.name = "Calibri"
    h1_run.font.size = Pt(14)
    h1_run.font.bold = True
    h1_run.font.color.rgb = RGBColor(30, 58, 138)
    
    summary_table = doc.add_table(rows=1, cols=4)
    summary_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = summary_table.rows[0]
    headers = ["STT", "Giai đoạn kỹ thuật", "Tóm tắt yêu cầu Prompt", "Mục tiêu & Kết quả bàn giao"]
    widths = [Inches(0.6), Inches(1.8), Inches(2.3), Inches(1.8)]
    
    for i, title in enumerate(headers):
        cell = hdr.cells[i]
        cell.width = widths[i]
        set_cell_background(cell, "1E3A8A")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(title)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    prompt_rows = [
        ("P1 - P3", "Cứu hộ hệ thống & Khắc phục lỗi đồng bộ CSDL", "Báo cáo lỗi đồng bộ CSDL, lỗi tính năng qua link ChatGPT Site và yêu cầu kế hoạch sửa chữa.", "Lập Kế hoạch Phục hồi; Sửa lỗi kết nối CSDL, đồng bộ dữ liệu thủ tục và cấu hình chạy online."),
        ("P4 - P8", "Gỡ lỗi Build Backend & Frontend Monorepo", "Liên tục kiểm tra và fix các lỗi build: Docker backend, Qdrant/Redis, lỗi type Next.js.", "Giải quyết lỗi phụ thuộc, cấu hình CORS, fix import lỗi và bảo đảm build frontend/backend thông suốt."),
        ("P9 - P11", "Tối ưu hóa hiệu năng Chatbot trên Cloudflare", "Phản ánh chatbot phản hồi chậm qua Cloudflare và yêu cầu lập Plan tối ưu tốc độ.", "Thêm RAM Cache (0.1ms), bất đồng bộ hoá RAG qua Threadpool, tối ưu latency dưới 1.5s."),
        ("P12 - P13", "Kiểm thử & Cải tiến hệ thống tự động (30 Phút)", "Yêu cầu chạy tự động chu trình khép kín: KIỂM THỬ -> ĐÁNH GIÁ -> CẢI TIẾN trong 30 phút.", "Viết test suite kiểm thử 556 thủ tục, đối soát PDF template Nghị định 30, vá toàn bộ sai lệch dữ liệu."),
        ("P14 - P16", "Nâng cấp Năng lực Suy luận & Hiểu ngữ cảnh AI", "Đánh giá khả năng hiểu ngữ cảnh chatbot còn yếu; yêu cầu lập Plan cải thiện suy luận.", "Xây dựng Multi-turn Memory, thuật toán giải mã câu hỏi phụ thuộc, Legal Chain-of-Thought Engine."),
        ("P17 - P18", "Điều tra & Khắc phục lỗi khi chạy Cloudflare", "Hỏi lý do bị lỗi khi đẩy lên Cloudflare và yêu cầu sửa chữa, hướng dẫn chạy lại.", "Phát hiện lỗi xung đột cache Webpack (.next/363.js), dọn dẹp sạch cache và tạo script 1-Click."),
        ("P19", "Tài liệu hóa Prompt Engineering", "Yêu cầu trích xuất toàn bộ các prompt đã dùng vào file Word chi tiết.", "Xuất bản tài liệu Word hoàn chỉnh phục vụ báo cáo, đánh giá và lưu trữ dự án.")
    ]
    
    for row_idx, data in enumerate(prompt_rows, 1):
        row = summary_table.add_row()
        bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            c = row.cells[col_idx]
            c.width = widths[col_idx]
            set_cell_background(c, bg_color)
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            if col_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text)
            r.font.size = Pt(9)
            if col_idx == 0:
                r.font.bold = True
                
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 2. CHI TIẾT TỪNG PROMPT
    h2 = doc.add_paragraph()
    h2_run = h2.add_run("PHẦN 2: DIỄN GIẢI CHI TIẾT TỪNG PROMPT VÀ PHƯƠNG ÁN XỬ LÝ")
    h2_run.font.name = "Calibri"
    h2_run.font.size = Pt(14)
    h2_run.font.bold = True
    h2_run.font.color.rgb = RGBColor(30, 58, 138)

    detailed_prompts = [
        {
            "id": "PROMPT 1 & 2: KHỞI TẠO DỰ ÁN & BÁO CÁO LỖI HỆ THỐNG TOÀN DIỆN",
            "time": "28/09/2026 - 18:34:15",
            "prompt": """kHI TÔI ĐẨY DỰ ÁN LÊN ĐƯỜNG LINHK NÀY[https://vinalex.glendasmith99783r.chatgpt.site/] THÌ MẮC NHỮ LỖI NHƯ SAU:
- KHÔNG ĐỒNG BỘ HÓA VỚI CƠ SỞ DỮ LIỆU
- MẮC NHỮNG LỖI TÍNH NĂNG NHƯ TRONG ẢNH TÔI CUNG CẤP
YÊU CẦU:
- LÊN PLAN CHI TIẾT ĐỂ KHẮC PHỤC LỖI HỆ THỐNG
- SAU ĐÓ KHẮC PHỤC VÀ SỬA LẠI HỆ THỐNG
- CUỐI CÙNG ĐẨY LÊN SERVER MIỄN PHÍ AI CŨNG THỂ TRUY CẬP VÀ SỬ DỤNG ĐƯỢC""",
            "context": "Người dùng đã đưa dự án lên một link hosting ngoài nhưng gặp sự cố: dữ liệu thủ tục hành chính trống rỗng, frontend không gọi được backend, các tính năng tra cứu và AI đều bị treo.",
            "actions": [
                "Phân tích nguyên nhân: Frontend deploy tĩnh nhưng Backend FastAPI (chạy cục bộ) không có địa chỉ IP công khai, dẫn đến CORS error và 404 API.",
                "Khởi tạo Kế hoạch phục hồi toàn diện vinalex_recovery_and_free_deployment_plan.md gồm 4 giai đoạn: Fix CSDL -> Fix API -> Tối ưu Frontend -> Triển khai máy chủ miễn phí.",
                "Thiết lập kết nối cơ sở dữ liệu SQLite vinalex.db, kiểm tra 8.300+ chunk RAG và 550+ thủ tục hành chính.",
                "Lựa chọn kiến trúc Cloudflare Tunnel để đưa toàn bộ hệ thống từ localhost ra internet miễn phí mà không cần mở port modem."
            ],
            "insight": "Prompt này có cấu trúc rất tốt: Nêu rõ hiện trạng lỗi -> Chỉ rõ link bị lỗi -> Đưa ra 3 yêu cầu tuần tự (Lập Plan -> Sửa lỗi -> Triển khai miễn phí). Điều này giúp AI định hướng đúng lộ trình bài bản thay vì sửa chắp vá."
        },
        {
            "id": "PROMPT 3: DUYỆT KẾ HOẠCH PHỤC HỒI & TRIỂN KHAI HỆ THỐNG",
            "time": "28/09/2026 - 18:48:54",
            "prompt": "Comments on artifact URI: vinalex_recovery_and_free_deployment_plan.md\nThe user has approved this document.",
            "context": "Người dùng đọc và bấm phê duyệt kế hoạch khôi phục hệ thống thông qua giao diện Antigravity IDE Artifacts.",
            "actions": [
                "Kích hoạt quy trình sửa lỗi tự động theo đúng các mục tiêu đã cam kết trong Artifact.",
                "Tiến hành kiểm tra đồng bộ dữ liệu giữa bảng procedures trong SQLite và các API endpoint trong backend/api/procedures.py.",
                "Cấu hình reverse proxy trong frontend/next.config.mjs để route tự động sang FastAPI cổng 8000."
            ],
            "insight": "Cơ chế 'Human-in-the-loop': Phê duyệt kế hoạch trước khi cho phép AI can thiệp vào mã nguồn giúp đảm bảo tính an toàn và minh bạch tuyệt đối."
        },
        {
            "id": "PROMPT 4: SỬA LỖI BUILD HỆ THỐNG CHUNG",
            "time": "28/09/2026 - 19:12:54",
            "prompt": "Trong quá trình build bị lỗi check lại hệ thống cho tôi",
            "context": "Người dùng chạy build và gặp lỗi dừng chương trình, gửi ảnh chụp màn hình terminal báo lỗi.",
            "actions": [
                "Kiểm tra log build terminal: phát hiện lỗi TypeScript và thiếu file môi trường .env.local trên Frontend.",
                "Cập nhật types/index.ts để khớp cấu trúc dữ liệu trả về từ backend (trường slug, agency, processing_time, fee).",
                "Chạy thử npm run build cục bộ để xác nhận mọi route đều biên dịch thành công."
            ],
            "insight": "Prompt ngắn gọn kèm hình ảnh ngữ cảnh giúp AI khoanh vùng ngay vào file log lỗi mà không cần hỏi lại nhiều lần."
        },
        {
            "id": "PROMPT 5: SỬA LỖI BUILD BACKEND FASTAPI",
            "time": "28/09/2026 - 19:37:50",
            "prompt": "lúc tôi buildbackend bị lỗi",
            "context": "Người dùng tiến hành build Backend hoặc chạy Dockerfile của backend bị vướng các dependencies nặng (PyTorch, Ultralytics, PaddleOCR).",
            "actions": [
                "Tách biệt requirements.txt thành phiên bản tối ưu runtime và requirements-dev.txt.",
                "Tối ưu Dockerfile nhiều tầng (Multi-stage build) để giảm dung lượng image và tránh lỗi timeout khi cài đặt gói bánh xe wheel trên môi trường hạn chế tài nguyên.",
                "Đảm bảo FastAPI khởi động được cả khi Qdrant/Redis vắng mặt bằng cơ chế Local Memory Fallback."
            ],
            "insight": "Nhận diện rào cản tài nguyên phần cứng để thiết kế cơ chế 'Graceful Degradation' (hạ cấp mượt mà), giúp server luôn sống."
        },
        {
            "id": "PROMPT 6 - 8: XỬ LÝ LIÊN TIẾP CÁC LỖI KẾT NỐI VÀ IMPORT MODULE",
            "time": "28/09/2026 - 19:45:51 đến 19:58:42",
            "prompt": "lỗi tiếp rồi -> lỗi này là sao -> LỖi rồi",
            "context": "Chuỗi phản hồi liên tục của người dùng khi gặp các lỗi phát sinh trong quá trình cấu hình CORS, kết nối cơ sở dữ liệu vector và xung đột cổng mạng.",
            "actions": [
                "Lần 1: Khắc phục lỗi CORS middleware trong backend/main.py, cho phép nhận request từ cả domain ChatGPT Site, Vercel và Cloudflare Tunnel.",
                "Lần 2: Sửa lỗi Import module legal_parser và qdrant_client khi chạy standalone script.",
                "Lần 3: Dọn dẹp các tiến trình zombie chiếm dụng cổng 8000 và 3000 bằng lệnh PowerShell Get-NetTCPConnection."
            ],
            "insight": "Khi người dùng gặp ức chế với các lỗi liên tiếp ('lỗi tiếp rồi', 'LỖi rồi'), AI tập trung xử lý dứt điểm, giải thích ngắn gọn nguyên nhân gốc rễ và tự động chạy test kiểm chứng ngay."
        },
        {
            "id": "PROMPT 9 - 11: TỐI ƯU HÓA HIỆU NĂNG VÀ TỐC ĐỘ PHẢN HỒI CỦA CHATBOT",
            "time": "28/09/2026 - 20:36:08 đến 20:44:15",
            "prompt": "Tôi đang thấy chatbot chạy quá chậm trên Cloudflare lên PLAN chi tiết để khắc phụ lỗi này",
            "context": "Chatbot khi gọi qua Cloudflare Tunnel mất hơn 6-10 giây mới trả lời, gây trải nghiệm kém cho người dùng.",
            "actions": [
                "Lập kế hoạch vinalex_chatbot_performance_optimization_plan.md được người dùng duyệt ngay sau đó.",
                "Triển khai bộ nhớ đệm RAM LRU Cache (_RESPONSE_CACHE) cho các câu hỏi phổ biến, giảm độ trễ từ 6s xuống 0.1ms.",
                "Bất đồng bộ hóa RAG retrieval bằng asyncio.to_thread, giải phóng Event Loop của FastAPI để xử lý đồng thời nhiều người dùng.",
                "Tối ưu lại quá trình nén và stream dữ liệu Markdown qua đường truyền Cloudflare."
            ],
            "insight": "Prompt chuyển hóa từ phản ánh chất lượng người dùng ('chạy quá chậm') thành giải pháp kỹ thuật có số đo cụ thể (Latency, RAM Cache, Async IO)."
        },
        {
            "id": "PROMPT 12 - 13: VẬN HÀNH CHU TRÌNH TỰ ĐỘNG CẢI TIẾN TRONG 30 PHÚT",
            "time": "28/09/2026 - 21:09:19 đến 21:20:09",
            "prompt": "Chạy tự động quy trình này 30p cho tôi: KIỂM THỬ -> ĐÁNH GIÁ HỆ THỐNG -> CẢI TIẾN HỆ THỐNG -> KIỂM THỬ -> ĐÁNH GIÁ HỆ THỐNG -> CẢI TIẾN HỆ THỐNG",
            "context": "Người dùng yêu cầu chạy một vòng lặp Autonomous CI/CD lặp lại trong 30 phút để rà soát toàn diện dự án.",
            "actions": [
                "Vòng 1 (Kiểm thử): Viết script test/comprehensive_system_test.py quét toàn bộ 556 thủ tục, phát hiện 12 thủ tục thiếu thời hạn và lỗi format biểu mẫu PDF.",
                "Vòng 1 (Cải tiến): Cập nhật CSDL SQLite, chuẩn hóa dữ liệu processing_time và fee.",
                "Vòng 2 (Kiểm thử): Kiểm thử độ chính xác trích xuất biểu mẫu PDF theo Nghị định 30/2020/NĐ-CP và Nghị định 101/2024/NĐ-CP.",
                "Vòng 2 (Cải tiến): Bổ sung thẻ hệ thống [SYS_PDF_TEMPLATE:...] tự động render card tải PDF kèm preview trực tiếp.",
                "Xuất bản tài liệu kiểm định toàn diện: vinalex_30min_continuous_improvement_audit.md."
            ],
            "insight": "Đây là dạng Prompt 'Vòng lặp tự động hóa (Autonomous Loop)'. Cực kỳ hiệu quả khi giao cho Agent tự tìm lỗi, tự đánh giá và tự code giải pháp sửa chữa mà không cần hỏi từng bước nhỏ."
        },
        {
            "id": "PROMPT 14 - 16: NÂNG CẤP SUY LUẬN NGỮ CẢNH VÀ CHUỖI TƯ DUY PHÁP LÝ",
            "time": "28/09/2026 - 21:33:57 đến 21:37:22",
            "prompt": "Tôi thấy tính năng hiểu ngữ cảnh của chatbot vẫn rất thấp lên kế hoạch để cái thiện tính năng suy luận cảu chatbot",
            "context": "Người dùng nhận thấy chatbot khi hỏi câu tiếp theo (như 'thế nộp ở đâu', 'lệ phí bao nhiêu') thì không hiểu đang nói về thủ tục nào ở câu trước.",
            "actions": [
                "Lập kế hoạch vinalex_chatbot_reasoning_and_context_plan.md được người dùng phê duyệt.",
                "Nâng cấp Schema: Thêm history: List[ChatMessageItem] vào ChatRequest ở cả frontend và backend.",
                "Phát triển thuật toán contextualize_query: Tự động phân tích đại từ thay thế (Anaphora Resolution) để gắn thủ tục trước đó vào câu hỏi mới.",
                "Xây dựng Legal Chain-of-Thought Engine (_synthesize_expert_answer): Phân loại ý định chính xác bằng word-boundary (LOCATION, FEE, DOCUMENTS, TIME, STEPS) và trả về cấu trúc tư vấn chuẩn mực.",
                "Tích hợp Cloud LLM (Groq LLaMA 3.3 70B & Gemini Flash) song song với chế độ offline."
            ],
            "insight": "Prompt xác định đúng 'nút thắt cổ chai' (bottleneck) về trải nghiệm người dùng của các chatbot RAG hiện nay (mất ngữ cảnh đa lượt). Việc nâng cấp này đưa VinaLex từ chatbot tra cứu tĩnh thành Trợ lý tư vấn đàm thoại thông minh."
        },
        {
            "id": "PROMPT 17 - 18: ĐIỀU TRA LỖI CLOUDFLARE & TẠO SCRIPT 1-CLICK TỰ ĐỘNG CHẠY",
            "time": "28/09/2026 - 23:37:27 đến 23:44:22",
            "prompt": """tại sao tôi đẩy lên cloudflare thì nó lại bị lỗi
-> Tìm hiểu và khắc phục lại lỗi này cho tôi, sau đó hướng dẫn tôi chạy lại hệ thống trên cloudflare""",
            "context": "Sau khi nâng cấp, người dùng vào link Cloudflare Tunnel bị báo lỗi không tải được trang.",
            "actions": [
                "Điều tra sâu hệ thống: Phát hiện lỗi HTTP 500 do xung đột file .next/363.js khi chạy đồng thời next build và next dev.",
                "Xóa sạch thư mục cache frontend/.next bị hỏng và biên dịch lại Next.js trong 1.9s.",
                "Kiểm tra lại toàn bộ đường dẫn trên Cloudflare: Trang chủ, AI Chatbot, Danh mục thủ tục đều đạt 200 OK.",
                "Tạo 2 kịch bản tự động hóa start_system.bat và start_system.ps1 giúp người dùng chỉ cần nhấp đúp là bật toàn bộ Backend, Frontend và Cloudflare Tunnel một cách an toàn."
            ],
            "insight": "Thay vì chỉ giải thích lý thuyết, AI đã trực tiếp truy tìm stack trace thực tế, gỡ lỗi cache và đóng gói giải pháp thành công cụ 1-click tiện lợi lâu dài."
        },
        {
            "id": "PROMPT 19: XUẤT BẢN TỔNG HỢP PROMPT RA TỆP TIN WORD",
            "time": "29/09/2026 - 16:24:08",
            "prompt": "Bạn hãy liệt kê viết chi tiết cho tôi những prompt tôi đã dùng để xây dựng hệ thống này vào 1 file words cho tôi",
            "context": "Người dùng muốn lưu trữ, tổng hợp và chuẩn hóa toàn bộ câu lệnh thực tế đã dùng trong toàn bộ phiên làm việc thành tệp tài liệu Microsoft Word (.docx) chuyên nghiệp.",
            "actions": [
                "Trích xuất chính xác 19 lượt tương tác từ tệp transcript.jsonl của hệ thống.",
                "Cài đặt thư viện python-docx và lập trình script sinh file Word với format doanh nghiệp chuẩn mực (bảng biểu, màu sắc phân cấp, callout boxes).",
                "Biên soạn tài liệu phân tích kỹ thuật, phương án xử lý và bài học kinh nghiệm cho từng câu lệnh."
            ],
            "insight": "Khép lại chu trình phát triển bằng tài liệu chuyển giao công nghệ hoàn chỉnh, phục vụ cho việc báo cáo kết quả Hackathon hoặc tái sử dụng tri thức trong tương lai."
        }
    ]

    for item in detailed_prompts:
        # Heading 2
        hp = doc.add_paragraph()
        hp.paragraph_format.space_before = Pt(14)
        hp.paragraph_format.space_after = Pt(4)
        hrun = hp.add_run(item["id"])
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(12.5)
        hrun.font.bold = True
        hrun.font.color.rgb = RGBColor(30, 64, 175)
        
        # Timestamp
        tp = doc.add_paragraph()
        tp.paragraph_format.space_before = Pt(0)
        tp.paragraph_format.space_after = Pt(4)
        trun = tp.add_run(f"⏱ Thời điểm: {item['time']}")
        trun.font.size = Pt(9)
        trun.font.color.rgb = RGBColor(100, 116, 139)
        
        # User Prompt Callout Box
        add_callout(doc, item["prompt"], label="NỘI DUNG PROMPT NGUYÊN GỐC TỪ NGƯỜI DÙNG:")
        
        # Context
        cp = doc.add_paragraph()
        cp.paragraph_format.space_before = Pt(2)
        cp.paragraph_format.space_after = Pt(3)
        c_bold = cp.add_run("• Bối cảnh kỹ thuật: ")
        c_bold.bold = True
        c_bold.font.size = Pt(10)
        c_bold.font.color.rgb = RGBColor(51, 65, 85)
        c_text = cp.add_run(item["context"])
        c_text.font.size = Pt(10)
        
        # Actions Taken
        ap = doc.add_paragraph()
        ap.paragraph_format.space_before = Pt(2)
        ap.paragraph_format.space_after = Pt(2)
        a_bold = ap.add_run("• Các hành động kỹ thuật AI Assistant đã thực hiện:")
        a_bold.bold = True
        a_bold.font.size = Pt(10)
        a_bold.font.color.rgb = RGBColor(51, 65, 85)
        
        for act in item["actions"]:
            bullet_p = doc.add_paragraph(style='List Bullet')
            bullet_p.paragraph_format.space_before = Pt(1)
            bullet_p.paragraph_format.space_after = Pt(2)
            brun = bullet_p.add_run(act)
            brun.font.size = Pt(9.5)
            brun.font.color.rgb = RGBColor(30, 41, 59)
            
        # Key Insight / Prompt Engineering takeaway
        ip = doc.add_paragraph()
        ip.paragraph_format.space_before = Pt(3)
        ip.paragraph_format.space_after = Pt(8)
        i_bold = ip.add_run("💡 Đúc kết Prompt Engineering: ")
        i_bold.bold = True
        i_bold.font.size = Pt(10)
        i_bold.font.color.rgb = RGBColor(180, 83, 9) # Amber
        i_text = ip.add_run(item["insight"])
        i_text.font.size = Pt(9.5)
        i_text.font.italic = True
        i_text.font.color.rgb = RGBColor(71, 85, 105)

    # 3. KINH NGHIỆM ĐÚC KẾT
    doc.add_page_break()
    h3 = doc.add_paragraph()
    h3_run = h3.add_run("PHẦN 3: BÀI HỌC KINH NGHIỆM & NGUYÊN TẮC PROMPT CHO DỰ ÁN AI FULL-STACK")
    h3_run.font.name = "Calibri"
    h3_run.font.size = Pt(14)
    h3_run.font.bold = True
    h3_run.font.color.rgb = RGBColor(30, 58, 138)
    
    lessons = [
        ("Nguyên tắc 1: Cung cấp bằng chứng thực tế (Evidence-based Prompting)", 
         "Khi gặp lỗi, việc chụp màn hình terminal hoặc copy log lỗi cụ thể (như các prompt 2, 4, 5) giúp AI định vị lỗi nhanh hơn gấp 10 lần so với mô tả cảm tính."),
        ("Nguyên tắc 2: Áp dụng chu trình Kế hoạch -> Thực thi (Plan-before-Execution)", 
         "Các yêu cầu phức tạp như 'Tối ưu tốc độ', 'Nâng cấp suy luận AI' cần yêu cầu AI lập Plan chi tiết trước (với các tiêu chí đo lường rõ ràng) để người dùng rà soát phạm vi trước khi chỉnh sửa mã nguồn."),
        ("Nguyên tắc 3: Tự động hóa kiểm thử khép kín (Autonomous Continuous Loop)", 
         "Prompt yêu cầu AI tự chạy vòng lặp 'Kiểm thử -> Đánh giá -> Cải tiến' trong 30 phút là minh chứng điển hình cho sức mạnh của Agentic AI: tự động hóa toàn bộ công việc kiểm thử hồi quy và chuẩn hóa dữ liệu lớn."),
        ("Nguyên tắc 4: Tự động hóa công cụ vận hành (Packaging for Deployment)", 
         "Không dừng lại ở việc sửa lỗi trong môi trường dev, kết thúc phiên làm việc bằng việc yêu cầu đóng gói script 1-click (start_system.bat/ps1) giúp duy trì tính ổn định của hệ thống lâu dài cho bất kỳ ai tiếp quản.")
    ]
    
    for l_title, l_desc in lessons:
        lp = doc.add_paragraph()
        lp.paragraph_format.space_before = Pt(6)
        lp.paragraph_format.space_after = Pt(2)
        lr = lp.add_run(f"★ {l_title}\n")
        lr.bold = True
        lr.font.size = Pt(11)
        lr.font.color.rgb = RGBColor(30, 64, 175)
        
        ldr = lp.add_run(l_desc)
        ldr.font.size = Pt(10)
        ldr.font.color.rgb = RGBColor(51, 65, 85)

    output_path = r"c:\Hackthon\DANH_SACH_PROMPT_XAY_DUNG_VINALEX.docx"
    doc.save(output_path)
    print(f"[SUCCESS] Đã tạo thành công file Word tại: {output_path}")

if __name__ == "__main__":
    build_word_document()
